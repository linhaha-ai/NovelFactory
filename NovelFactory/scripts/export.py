# -*- coding: utf-8 -*-
"""导出：合并全部章节为 单文件txt / html / docx（docx 需 python-docx，缺省时跳过）。"""
from __future__ import annotations

import argparse
from html import escape
from pathlib import Path

from ..core.utils import read_text

DEFAULT_WS = Path(__file__).resolve().parent.parent.parent / "novel_factory"


def collect_chapters(workspace: Path):
    ms = workspace / "manuscript"
    if not ms.exists():
        return []
    chapters = sorted(ms.glob("ch*.md"), key=lambda p: p.name)
    return [(p.name, read_text(p)) for p in chapters]


def export_txt(workspace: Path, out_dir: Path) -> Path:
    out_dir.mkdir(parents=True, exist_ok=True)
    title = "未命名"
    cfg = workspace / "config.toml"
    if cfg.exists():
        for ln in cfg.read_text(encoding="utf-8").splitlines():
            if ln.strip().startswith("title"):
                title = ln.split("=", 1)[1].strip().strip('"').strip("'")
                break
    out = out_dir / f"{title}_合集.txt"
    parts = [f"《{title}》\n"]
    for name, content in collect_chapters(workspace):
        parts.append(f"\n{'=' * 20}\n{name}\n{'=' * 20}\n\n{content}\n")
    out.write_text("\n".join(parts), encoding="utf-8")
    return out


def export_html(workspace: Path, out_dir: Path) -> Path:
    out_dir.mkdir(parents=True, exist_ok=True)
    title = Path(workspace).name
    body = []
    for name, content in collect_chapters(workspace):
        body.append(f"<h2>{escape(name)}</h2>\n"
                    f"<div style='line-height:1.9;white-space:pre-wrap;'>{escape(content)}</div>")
    html = f"<!DOCTYPE html><html lang='zh'><meta charset='utf-8'><title>{escape(title)}</title>" \
           f"<body style='max-width:760px;margin:0 auto;font-family:serif;'><h1>{escape(title)}</h1>" \
           f"{''.join(body)}</body></html>"
    out = out_dir / f"{title}.html"
    out.write_text(html, encoding="utf-8")
    return out


def export_docx(workspace: Path, out_dir: Path) -> Path | None:
    try:
        import docx  # type: ignore
    except ImportError:
        print("⚠ 未安装 python-docx，跳过 docx 导出（可执行 pip install python-docx）")
        return None
    out_dir.mkdir(parents=True, exist_ok=True)
    title = Path(workspace).name
    doc = docx.Document()
    doc.add_heading(title, level=0)
    for name, content in collect_chapters(workspace):
        doc.add_heading(name, level=1)
        for para in content.splitlines():
            if para.strip():
                doc.add_paragraph(para)
    out = out_dir / f"{title}.docx"
    doc.save(str(out))
    return out


def export_all(workspace: str | Path, formats: str = "txt,html",
               out_dir: str | Path | None = None) -> list:
    ws, out_dir = Path(workspace), Path(out_dir or Path(workspace) / "export")
    done = []
    for fmt in [f.strip() for f in formats.split(",") if f.strip()]:
        if fmt == "txt":
            p = export_txt(ws, out_dir)
            print(f"✅ {p}")
            done.append(p)
        elif fmt == "html":
            p = export_html(ws, out_dir)
            print(f"✅ {p}")
            done.append(p)
        elif fmt == "docx":
            p = export_docx(ws, out_dir)
            if p:
                print(f"✅ {p}")
                done.append(p)
        else:
            print(f"⚠ 未知格式: {fmt}")
    return done


def main() -> None:
    ap = argparse.ArgumentParser(description="导出小说合集")
    ap.add_argument("--workspace", default=str(DEFAULT_WS))
    ap.add_argument("--out", default=str(DEFAULT_WS / "export"))
    ap.add_argument("--formats", default="txt,html", help="txt,html,docx")
    args = ap.parse_args()
    export_all(args.workspace, args.formats, args.out)


if __name__ == "__main__":
    main()
